import json
import sys

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter


def safe_title(title, existing):
    title = (
        title.replace("\\", "＼")
        .replace("/", "／")
        .replace("?", "-")
        .replace("*", "-")
        .replace("[", "(")
        .replace("]", ")")
    )[:31]

    base = title
    counter = 2

    while title in existing:
        suffix = f" {counter}"
        title = (base[: 31 - len(suffix)] + suffix)[:31]
        counter += 1

    existing.add(title)
    return title


def write_workbook(input_path, output_path):
    with open(input_path, "r", encoding="utf-8") as handle:
        sheets = json.load(handle)

    workbook = Workbook()
    workbook.remove(workbook.active)
    existing_titles = set()

    for sheet in sheets:
        worksheet = workbook.create_sheet(
            title=safe_title(sheet["name"], existing_titles)
        )

        rows = sheet.get("rows", [])

        for row in rows:
            worksheet.append(row)

        if rows:
            for cell in worksheet[1]:
                cell.font = Font(bold=True)

        worksheet.freeze_panes = "A2"

        for column in worksheet.columns:
            width = 12

            for cell in column:
                value = "" if cell.value is None else str(cell.value)
                width = max(width, min(len(value) + 2, 40))

            worksheet.column_dimensions[get_column_letter(column[0].column)].width = width

    workbook.save(output_path)


if __name__ == "__main__":
    write_workbook(sys.argv[1], sys.argv[2])
